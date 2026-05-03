from pathlib import Path

from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw, ImageFont
from copy import deepcopy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE_PDF = Path(r"C:\Users\gokul\Downloads\Ex-05.pdf")
GENERATED_CODE_PDF = ROOT / "Ex-05-code-pages.pdf"
GENERATED_OVERLAY_PDF = ROOT / "Ex-05-code-overlay.pdf"
GENERATED_RESULT_PDF = ROOT / "Ex-05-result-overlay.pdf"
OUTPUT_PDF = ROOT / "Ex-05-updated.pdf"
OUTPUT_DOCX = ROOT / "Ex-05-updated.docx"
RESULT_IMAGES = [
    ROOT / "result_screenshot_1.png",
    ROOT / "result_screenshot_2.png",
    ROOT / "result_screenshot_3.png",
]
RESULT_COLLAGE = ROOT / "result_screenshots_collage.png"

PAGE_WIDTH, PAGE_HEIGHT = letter
LEFT = 58
RIGHT = 52
TOP = 92
BOTTOM = 56
CODE_FONT = "Courier"
CODE_SIZE = 8.7
LINE_HEIGHT = 10.9
MAX_WIDTH = PAGE_WIDTH - LEFT - RIGHT
CONTENT_TOP = 682
CONTENT_BOTTOM = 58


def read(relative_path: str) -> str:
    return (ROOT / relative_path).read_text(encoding="utf-8").rstrip()


def relevant_version_catalog() -> str:
    text = read("gradle/libs.versions.toml")
    keep = []
    wanted_prefixes = (
        "agp",
        "kotlin",
        "composeBom",
        "mpandroidchart",
        "activityCompose",
        "coreKtx",
        "lifecycleRuntimeKtx",
    )
    wanted_libraries = (
        "mpandroidchart",
        "androidx-core-ktx",
        "androidx-lifecycle-runtime-ktx",
        "androidx-activity-compose",
        "androidx-compose-bom",
        "androidx-compose-ui",
        "androidx-compose-ui-graphics",
        "androidx-compose-ui-tooling-preview",
        "androidx-compose-material3",
    )
    wanted_plugins = ("android-application", "kotlin-compose")
    section = None

    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("[") and stripped.endswith("]"):
            section = stripped
            keep.append(line)
            continue
        if not stripped:
            continue
        key = stripped.split("=", 1)[0].strip()
        if section == "[versions]" and key in wanted_prefixes:
            keep.append(line)
        elif section == "[libraries]" and key in wanted_libraries:
            keep.append(line)
        elif section == "[plugins]" and key in wanted_plugins:
            keep.append(line)
    return "\n".join(keep)


SECTIONS = [
    ("AndroidManifest.xml", read("app/src/main/AndroidManifest.xml")),
    ("build.gradle.kts (Module: app)", read("app/build.gradle.kts")),
    ("libs.versions.toml (Required Entries)", relevant_version_catalog()),
    ("MainActivity.kt", read("app/src/main/java/com/example/exp5/MainActivity.kt")),
    ("Color.kt", read("app/src/main/java/com/example/exp5/ui/theme/Color.kt")),
    ("Theme.kt", read("app/src/main/java/com/example/exp5/ui/theme/Theme.kt")),
    ("Type.kt", read("app/src/main/java/com/example/exp5/ui/theme/Type.kt")),
    ("strings.xml", read("app/src/main/res/values/strings.xml")),
    ("themes.xml", read("app/src/main/res/values/themes.xml")),
]

STUDENT_NAME = "N GOKUL KRISHNA"
REGISTER_NO = "231701014"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def rounded(draw: ImageDraw.ImageDraw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def draw_phone_shell(draw: ImageDraw.ImageDraw, w: int, h: int) -> None:
    rounded(draw, (1, 1, w - 2, h - 2), 35, "#151515", "#4b4b4b", 2)
    rounded(draw, (7, 7, w - 8, h - 8), 31, "#252525", "#0a0a0a", 2)
    rounded(draw, (22, 24, w - 22, h - 26), 18, "#0f172a")
    draw.ellipse((190, 30, 215, 55), fill="#101820", outline="#5f6b7a", width=2)
    draw.ellipse((199, 38, 206, 45), fill="#8ba2bb")
    draw.rectangle((2, 108, 6, 153), fill="#2a2a2a")
    draw.rectangle((w - 6, 154, w - 2, 220), fill="#2a2a2a")


def draw_status(draw: ImageDraw.ImageDraw) -> None:
    small = font(11)
    draw.text((42, 39), "1:25", fill="#0b1220", font=small)
    for x in [86, 107, 128, 149, 170]:
        draw.ellipse((x, 42, x + 8, 50), outline="#0b1220", width=1)
    draw.polygon([(322, 41), (335, 41), (328, 49)], fill="#0b1220")
    draw.rectangle((344, 40, 351, 50), fill="#0b1220")
    draw.rectangle((356, 39, 364, 50), fill="#0b1220")


def draw_header(draw: ImageDraw.ImageDraw) -> None:
    draw.text((40, 83), "Financial Overview", fill="white", font=font(24, True))
    draw.text(
        (40, 121),
        "Your revenue performance over the last 6 months",
        fill="#9ca3af",
        font=font(13),
    )


def draw_summary_cards(draw: ImageDraw.ImageDraw) -> None:
    for x, title, amount, trend, color in [
        (39, "Total Balance", "$45,231.89", "+12.5%", "#60a5fa"),
        (214, "Monthly Profit", "$8,432.00", "+5.2%", "#2dd4bf"),
    ]:
        rounded(draw, (x, 159, x + 162, 259), 12, "#1e293b")
        draw.text((x + 15, 179), title, fill="#94a3b8", font=font(11))
        draw.text((x + 15, 200), amount, fill="white", font=font(18, True))
        draw.text((x + 15, 229), trend, fill=color, font=font(11))


def chart_panel(draw: ImageDraw.ImageDraw, box, title):
    rounded(draw, box, 22, "#1e293b")
    draw.text((box[0] + 15, box[1] + 18), title, fill="white", font=font(16, True))


def draw_line_chart(draw: ImageDraw.ImageDraw, top: int, compact: bool = False) -> None:
    chart_panel(draw, (39, top, 377, top + 312), "Revenue Stream")
    gx0, gy0, gx1, gy1 = 88, top + 84, 350, top + 265
    vals = [12000, 15000, 13500, 19000, 17500, 24000]
    labels = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
    for i, yv in enumerate([12000, 14000, 16000, 18000, 20000, 22000, 24000]):
        y = gy1 - (yv - 12000) / 12000 * (gy1 - gy0)
        draw.line((gx0, y, gx1, y), fill="#334155", width=1)
        draw.text((56, y - 6), f"{yv:,}", fill="white", font=font(8))
    points = []
    for i, v in enumerate(vals):
        x = gx0 + i * (gx1 - gx0) / 5
        y = gy1 - (v - 12000) / 12000 * (gy1 - gy0)
        points.append((x, y))
    fill_poly = points + [(gx1, gy1), (gx0, gy1)]
    draw.polygon(fill_poly, fill="#233f70")
    draw.line(points, fill="#3b82f6", width=2, joint="curve")
    for x, y in points:
        draw.ellipse((x - 2, y - 2, x + 2, y + 2), fill="#3b82f6")
    for i, lab in enumerate(labels):
        x = gx0 + i * (gx1 - gx0) / 5
        draw.text((x - 10, gy1 + 8), lab, fill="white", font=font(8))


def draw_bar_chart(draw: ImageDraw.ImageDraw, top: int, panel_height: int = 312) -> None:
    chart_panel(draw, (39, top, 377, top + panel_height), "Monthly Comparison")
    gx0, gy0, gx1, gy1 = 88, top + 94, 343, top + panel_height - 50
    vals = [15000, 18000, 16000, 22000, 20000, 28000]
    for yv in [15000, 18000, 21000, 24000, 27000]:
        y = gy1 - (yv - 14000) / 14000 * (gy1 - gy0)
        draw.line((gx0, y, gx1, y), fill="#334155", width=1)
        draw.text((52, y - 6), f"{yv:,}", fill="white", font=font(8))
    bar_w = 26
    for i, v in enumerate(vals):
        x = gx0 + 16 + i * 43
        y = gy1 - (v - 14000) / 14000 * (gy1 - gy0)
        draw.rectangle((x, y, x + bar_w, gy1), fill="#16c28d")
        draw.text((x + 4, gy1 + 6), str(i), fill="white", font=font(8))
        draw.text((x + 3, y - 9), f"{v:,}", fill="white", font=font(5))


def draw_pie_chart(draw: ImageDraw.ImageDraw, top: int) -> None:
    chart_panel(draw, (32, top, 370, top + 312), "Expense Categories")
    cx, cy, r = 158, top + 190, 105
    vals = [40, 25, 15, 10, 10]
    labels = ["Housing", "Food", "Transport", "Entertainment", "Other"]
    cols = ["#3b82f6", "#10b981", "#f59e0b", "#ef4444", "#8b5cf6"]
    start = -90
    total = sum(vals)
    for val, col, lab in zip(vals, cols, labels):
        end = start + val / total * 360
        draw.pieslice((cx - r, cy - r, cx + r, cy + r), start, end, fill=col)
        start = end
    draw.ellipse((cx - 55, cy - 55, cx + 55, cy + 55), fill="#1e293b")
    lx, ly = 278, top + 154
    for i, (col, lab) in enumerate(zip(cols, labels)):
        y = ly + i * 14
        draw.rectangle((lx, y, lx + 8, y + 8), fill=col)
        draw.text((lx + 13, y - 2), lab, fill="white", font=font(8))
    draw.text((lx + 13, ly + 72), "Categories", fill="white", font=font(7))


def draw_home_indicator(draw: ImageDraw.ImageDraw, y: int = 825) -> None:
    rounded(draw, (160, y, 250, y + 4), 2, "white")


def generate_result_images() -> None:
    for idx, path in enumerate(RESULT_IMAGES):
        img = Image.new("RGB", (413, 864), "#ffffff")
        d = ImageDraw.Draw(img)
        draw_phone_shell(d, 413, 864)
        draw_status(d)
        if idx in (0, 1):
            draw_header(d)
            draw_summary_cards(d)
            draw_line_chart(d, 280)
            draw_bar_chart(d, 607, 260)
            draw_home_indicator(d, 824)
        else:
            draw_line_chart(d, -190)
            draw_bar_chart(d, 137, 312)
            draw_pie_chart(d, 463)
            draw_home_indicator(d, 826)
        img.save(path)

    thumbs = []
    for path in RESULT_IMAGES:
        img = Image.open(path).convert("RGB")
        target_w = 320
        target_h = int(img.height * target_w / img.width)
        thumbs.append(img.resize((target_w, target_h), Image.Resampling.LANCZOS))
    gap = 28
    collage = Image.new("RGB", (target_w * 3 + gap * 2, max(t.height for t in thumbs)), "white")
    x = 0
    for thumb in thumbs:
        collage.paste(thumb, (x, 0))
        x += target_w + gap
    collage.save(RESULT_COLLAGE)


def split_code_line(line: str) -> list[str]:
    if not line:
        return [""]

    available = MAX_WIDTH
    result = []
    pending = line.replace("\t", "    ")
    max_chars = max(34, int(available / stringWidth("M", CODE_FONT, CODE_SIZE)))

    while stringWidth(pending, CODE_FONT, CODE_SIZE) > available:
        cut = min(len(pending), max_chars)
        while cut > 18 and stringWidth(pending[:cut], CODE_FONT, CODE_SIZE) > available:
            cut -= 1

        break_at = max(
            pending.rfind(" ", 0, cut),
            pending.rfind(",", 0, cut),
            pending.rfind(")", 0, cut),
            pending.rfind("}", 0, cut),
        )
        if break_at < 18:
            break_at = cut

        result.append(pending[:break_at].rstrip())
        pending = "    " + pending[break_at:].lstrip()

    result.append(pending)
    return result


def draw_title(c: canvas.Canvas, title: str, y: float) -> float:
    c.setFont("Times-Bold", 11)
    c.setFillColor(colors.black)
    c.drawString(LEFT, y, title)
    return y - 16


def draw_code_line(c: canvas.Canvas, line: str, y: float) -> None:
    c.setFont(CODE_FONT, CODE_SIZE)
    c.setFillColor(colors.black)
    c.drawString(LEFT, y, line)


def create_overlay_pdf() -> int:
    c = canvas.Canvas(str(GENERATED_OVERLAY_PDF), pagesize=letter)
    y = CONTENT_TOP

    def new_page():
        nonlocal y
        c.showPage()
        y = CONTENT_TOP

    for title, code in SECTIONS:
        if y < BOTTOM + 42:
            new_page()
        y = draw_title(c, title, y)

        for raw_line in code.splitlines():
            for line in split_code_line(raw_line):
                if y < BOTTOM:
                    new_page()
                draw_code_line(c, line, y)
                y -= LINE_HEIGHT
        y -= 12

    c.save()
    return len(PdfReader(str(GENERATED_OVERLAY_PDF)).pages)


def create_result_pdf() -> int:
    generate_result_images()
    c = canvas.Canvas(str(GENERATED_RESULT_PDF), pagesize=letter)

    # Page 55-style result page, with the old output images and duplicate labels covered.
    c.setFillColor(colors.white)
    c.rect(46, 76, 520, 624, fill=1, stroke=0)
    c.setFillColor(colors.black)
    c.setFont("Times-Bold", 12)
    c.drawString(58, 688, "Output")
    c.drawImage(str(RESULT_IMAGES[0]), 75, 238, width=150, height=314, preserveAspectRatio=True, mask="auto")
    c.drawImage(str(RESULT_IMAGES[1]), 238, 238, width=150, height=314, preserveAspectRatio=True, mask="auto")
    c.drawImage(str(RESULT_IMAGES[2]), 397, 238, width=145, height=304, preserveAspectRatio=True, mask="auto")
    c.setFont("Times-Bold", 12)
    c.drawString(58, 122, "Result")
    c.setFont("Times-Roman", 12)
    c.drawString(
        58,
        96,
        "Thus, the data visualization mobile application was designed, implemented, and executed successfully.",
    )
    c.save()
    return 1


def create_student_overlay() -> Path:
    path = ROOT / "Ex-05-student-overlay.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFillColor(colors.white)
    c.rect(171, 643, 336, 18, fill=1, stroke=0)
    c.setFillColor(colors.black)
    c.setFont("Times-Roman", 11)
    c.drawString(172, 649, REGISTER_NO)
    c.drawString(280, 649, f"Name: {STUDENT_NAME}")
    c.save()
    return path


def merge_pdf() -> None:
    generated_pages = create_overlay_pdf()
    create_result_pdf()
    student_overlay = PdfReader(str(create_student_overlay()))
    source = PdfReader(str(SOURCE_PDF))
    overlay = PdfReader(str(GENERATED_OVERLAY_PDF))
    result_overlay = PdfReader(str(GENERATED_RESULT_PDF))
    writer = PdfWriter()

    first_page = deepcopy(source.pages[0])
    first_page.merge_page(student_overlay.pages[0])
    writer.add_page(first_page)
    for i in range(1, 3):
        writer.add_page(deepcopy(source.pages[i]))

    for i, overlay_page in enumerate(overlay.pages):
        template_index = min(3 + i, len(source.pages) - 1)
        page = deepcopy(source.pages[template_index])

        c = canvas.Canvas(str(GENERATED_CODE_PDF), pagesize=letter)
        c.setFillColor(colors.white)
        c.rect(44, CONTENT_BOTTOM - 7, 524, CONTENT_TOP - CONTENT_BOTTOM + 8, fill=1, stroke=0)
        c.save()
        cover = PdfReader(str(GENERATED_CODE_PDF)).pages[0]
        page.merge_page(cover)
        page.merge_page(overlay_page)
        writer.add_page(page)

    result_page = deepcopy(source.pages[-1])
    result_page.merge_page(result_overlay.pages[0])
    writer.add_page(result_page)

    with OUTPUT_PDF.open("wb") as fh:
        writer.write(fh)

    print(f"Created {OUTPUT_PDF}")
    print(f"Total pages: {len(writer.pages)}; regenerated code pages: {generated_pages}")


def set_page_border(section) -> None:
    sect_pr = section._sectPr
    pg_borders = sect_pr.find(qn("w:pgBorders"))
    if pg_borders is None:
        pg_borders = OxmlElement("w:pgBorders")
        sect_pr.append(pg_borders)
    pg_borders.set(qn("w:offsetFrom"), "page")
    for side in ["top", "left", "bottom", "right"]:
        el = pg_borders.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            pg_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "12")
        el.set(qn("w:space"), "18")
        el.set(qn("w:color"), "000000")


def add_header(paragraph, page_no: int) -> None:
    run = paragraph.add_run(
        "Department of Computer Science and Design    |     Rajalakshmi Engineering College              ."
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    page = paragraph.insert_paragraph_before()
    page.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_run = page.add_run(str(page_no))
    page_run.font.name = "Times New Roman"
    page_run.font.size = Pt(10)


def add_code(doc: Document, title: str, code: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    for line in code.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line if line else " ")
        r.font.name = "Courier New"
        r.font.size = Pt(8.7)


def create_docx() -> None:
    generate_result_images()
    source = PdfReader(str(SOURCE_PDF))
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    set_page_border(section)

    body_text = "\n".join((source.pages[i].extract_text() or "") for i in range(3))
    lines = [line.strip() for line in body_text.splitlines() if line.strip()]
    filtered = []
    for line in lines:
        if line.startswith("Department of") or line in {"43", "44", "45"}:
            continue
        filtered.append(line)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(8)
    r = meta.add_run("Ex. No.  :  5     Date: \n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r = meta.add_run(f"Register No.: {REGISTER_NO}      Name: {STUDENT_NAME}")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)

    for line in filtered:
        if line in {"Ex. No.  :  5     Date:", "Register No.:      Name:"}:
            continue
        p = doc.add_paragraph()
        if line in {"Data Visualization App", "Aim", "Procedure"}:
            r = p.add_run(line)
            r.bold = True
            r.font.size = Pt(12 if line != "Data Visualization App" else 14)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if line == "Data Visualization App" else WD_ALIGN_PARAGRAPH.LEFT
        else:
            r = p.add_run(line)
            r.font.size = Pt(10.5)
        r.font.name = "Times New Roman"
        p.paragraph_format.space_after = Pt(3)

    doc.add_page_break()
    for title, code in SECTIONS:
        add_code(doc, title, code)

    doc.add_page_break()
    p = doc.add_paragraph()
    r = p.add_run("Output")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(RESULT_COLLAGE), width=Inches(6.2))
    p = doc.add_paragraph()
    r = p.add_run("Result")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    r = p.add_run(
        "Thus, the data visualization mobile application was designed, implemented, and executed successfully."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)

    for section in doc.sections:
        set_page_border(section)
    doc.save(OUTPUT_DOCX)
    print(f"Created {OUTPUT_DOCX}")


if __name__ == "__main__":
    merge_pdf()
    create_docx()
